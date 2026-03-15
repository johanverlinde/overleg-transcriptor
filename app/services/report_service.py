"""Service voor het genereren van Word-rapportages."""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Dict, List, Optional

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

if TYPE_CHECKING:
    from app.services.meeting_analysis_service import GemeenteraadAnalyse, GesprekAnalyse, TeamOverlegAnalyse


def _add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_heading(level=level)
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(14 if level == 1 else 12)


def _add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.style = "Normal"


def build_word_report(output_path: Path, analysis: Dict[str, object]) -> Path:
    """Legacy functie voor eenvoudige vergaderverslagen."""
    document = Document()

    _add_heading(document, "Verslag overleg (automatisch gegenereerd)", level=1)

    # 1. Basisgegevens
    _add_heading(document, "1. Basisgegevens", level=2)
    basis = analysis.get("basisgegevens", {})
    _add_paragraph(document, f"Datum: {basis.get('datum', 'Onbekend')}")
    _add_paragraph(document, f"Tijd: {basis.get('tijd', 'Onbekend')}")
    _add_paragraph(document, f"Locatie / online: {basis.get('locatie', 'Onbekend')}")
    _add_paragraph(document, f"Aanwezig: {basis.get('aanwezigen', 'Onbekend')}")

    # 2. Korte samenvatting
    _add_heading(document, "2. Korte samenvatting", level=2)
    _add_paragraph(document, analysis.get("summary", "Geen samenvatting beschikbaar."))

    # 3. Belangrijkste bespreekpunten
    _add_heading(document, "3. Belangrijkste bespreekpunten", level=2)
    for index, topic in enumerate(analysis.get("topics", []), start=1):
        _add_heading(document, f"3.{index} {topic['title']}", level=3)
        _add_paragraph(document, topic.get("details", ""))

    # 4. Besluiten en afspraken
    _add_heading(document, "4. Besluiten en afspraken", level=2)
    for decision in analysis.get("decisions", []):
        document.add_paragraph(
            f"Besluit: {decision.get('title', '')}", style="List Bullet"
        )
        document.add_paragraph(
            f"Toelichting: {decision.get('context', '')}", style="List Continue"
        )

    # 5. Actiepunten
    _add_heading(document, "5. Actiepunten", level=2)
    for action in analysis.get("actions", []):
        document.add_paragraph(
            f"Actie: {action.get('action', '')}", style="List Number"
        )
        document.add_paragraph(
            f"Eigenaar: {action.get('owner', 'Onbekend')}", style="List Continue"
        )
        document.add_paragraph(
            f"Deadline: {action.get('deadline', 'N.t.b.')}", style="List Continue"
        )
        document.add_paragraph(
            f"Opmerkingen: {action.get('notes', '')}", style="List Continue"
        )

    # 6. Risico's en aandachtspunten
    _add_heading(document, "6. Risico's en aandachtspunten", level=2)
    risks: List[Dict[str, str]] = analysis.get("risks", [])  # type: ignore[assignment]
    if risks:
        for risk in risks:
            document.add_paragraph(
                f"Risico: {risk.get('description', '')}", style="List Bullet"
            )
            document.add_paragraph(
                f"Maatregel: {risk.get('mitigation', '')}", style="List Continue"
            )
    else:
        _add_paragraph(document, "Geen risico's benoemd.")

    document.add_page_break()
    _add_heading(document, "Bijlage: opgeschoonde transcriptie", level=2)
    _add_paragraph(document, analysis.get("clean_transcript", "Transcriptie niet beschikbaar."))

    document.save(output_path)
    return output_path


def build_gemeenteraad_report(
    output_path: Path,
    analyse: "GemeenteraadAnalyse",
    metadata: Optional[Dict[str, str]] = None,
    transcript: Optional[str] = None,
) -> Path:
    """
    Genereer een uitgebreid Word-rapport voor een gemeenteraadsvergadering.
    
    Args:
        output_path: Waar het rapport opgeslagen wordt
        analyse: De GemeenteraadAnalyse met alle geëxtraheerde informatie
        metadata: Optionele metadata (datum, commissie, etc.)
        transcript: Optionele transcriptie voor bijlage
        
    Returns:
        Path naar het gegenereerde rapport
    """
    document = Document()
    metadata = metadata or {}
    
    # Titel
    title = document.add_heading(level=0)
    title_run = title.add_run("Verslag Gemeenteraadsvergadering")
    title_run.font.size = Pt(18)
    title_run.font.name = "Calibri"
    
    if metadata.get("commissie"):
        subtitle = document.add_paragraph()
        subtitle_run = subtitle.add_run(metadata["commissie"])
        subtitle_run.font.size = Pt(14)
        subtitle_run.font.italic = True
    
    # Automatisch gegenereerd label
    auto_label = document.add_paragraph()
    auto_label.add_run("Dit verslag is automatisch gegenereerd op basis van de transcriptie van de vergadering.").italic = True
    
    document.add_paragraph()  # Spacing
    
    # =========================================================================
    # 1. BASISGEGEVENS
    # =========================================================================
    _add_heading(document, "1. Basisgegevens", level=1)
    
    table = document.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    
    cells = [
        ("Datum", metadata.get("datum", "Niet gespecificeerd")),
        ("Tijd", metadata.get("tijd", "Niet gespecificeerd")),
        ("Locatie", metadata.get("locatie", "Niet gespecificeerd")),
        ("Voorzitter", metadata.get("voorzitter", "Niet gespecificeerd")),
    ]
    
    for i, (label, value) in enumerate(cells):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
    
    document.add_paragraph()
    
    # =========================================================================
    # 2. SAMENVATTING
    # =========================================================================
    _add_heading(document, "2. Samenvatting", level=1)
    _add_paragraph(document, analyse.samenvatting or "Geen samenvatting beschikbaar.")
    
    # Belangrijkste punten
    if analyse.belangrijkste_punten:
        _add_heading(document, "Belangrijkste punten", level=2)
        for punt in analyse.belangrijkste_punten:
            document.add_paragraph(punt, style="List Bullet")
    
    # =========================================================================
    # 3. MOTIES
    # =========================================================================
    _add_heading(document, "3. Moties", level=1)
    
    if analyse.moties:
        for i, motie in enumerate(analyse.moties, 1):
            _add_heading(document, f"3.{i} {motie.titel or 'Motie'}", level=2)
            
            if motie.nummer:
                _add_paragraph(document, f"Nummer: {motie.nummer}")
            
            indiener_text = motie.indiener
            if motie.partij:
                indiener_text += f" ({motie.partij})"
            if indiener_text:
                _add_paragraph(document, f"Indiener: {indiener_text}")
            
            if motie.samenvatting:
                _add_paragraph(document, f"Inhoud: {motie.samenvatting}")
            
            if motie.stemuitslag:
                stemuitslag_para = document.add_paragraph()
                stemuitslag_para.add_run("Stemuitslag: ").bold = True
                stemuitslag_para.add_run(motie.stemuitslag)
            
            document.add_paragraph()  # Spacing
    else:
        _add_paragraph(document, "Er zijn geen moties behandeld in deze vergadering.")
    
    # =========================================================================
    # 4. AMENDEMENTEN
    # =========================================================================
    _add_heading(document, "4. Amendementen", level=1)
    
    if analyse.amendementen:
        for i, amendement in enumerate(analyse.amendementen, 1):
            _add_heading(document, f"4.{i} {amendement.titel or 'Amendement'}", level=2)
            
            if amendement.nummer:
                _add_paragraph(document, f"Nummer: {amendement.nummer}")
            
            indiener_text = amendement.indiener
            if amendement.partij:
                indiener_text += f" ({amendement.partij})"
            if indiener_text:
                _add_paragraph(document, f"Indiener: {indiener_text}")
            
            if amendement.wijziging:
                _add_paragraph(document, f"Wijziging: {amendement.wijziging}")
            
            if amendement.stemuitslag:
                stemuitslag_para = document.add_paragraph()
                stemuitslag_para.add_run("Stemuitslag: ").bold = True
                stemuitslag_para.add_run(amendement.stemuitslag)
            
            document.add_paragraph()
    else:
        _add_paragraph(document, "Er zijn geen amendementen behandeld in deze vergadering.")
    
    # =========================================================================
    # 5. TOEZEGGINGEN
    # =========================================================================
    _add_heading(document, "5. Toezeggingen", level=1)
    
    if analyse.toezeggingen:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        
        # Header
        headers = ["Door", "Aan", "Toezegging", "Deadline"]
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        
        for toezegging in analyse.toezeggingen:
            row = table.add_row()
            row.cells[0].text = toezegging.door or "-"
            row.cells[1].text = toezegging.aan or "-"
            row.cells[2].text = toezegging.onderwerp or "-"
            row.cells[3].text = toezegging.deadline or "-"
        
        document.add_paragraph()
        
        # Details per toezegging
        for i, toezegging in enumerate(analyse.toezeggingen, 1):
            if toezegging.context:
                _add_heading(document, f"Toelichting toezegging {i}", level=2)
                _add_paragraph(document, toezegging.context)
    else:
        _add_paragraph(document, "Er zijn geen toezeggingen gedaan in deze vergadering.")
    
    # =========================================================================
    # 6. BESLUITEN
    # =========================================================================
    _add_heading(document, "6. Besluiten", level=1)
    
    if analyse.besluiten:
        for i, besluit in enumerate(analyse.besluiten, 1):
            titel = besluit.get("titel", f"Besluit {i}")
            _add_heading(document, f"6.{i} {titel}", level=2)
            
            if besluit.get("besluit"):
                _add_paragraph(document, besluit["besluit"])
            
            if besluit.get("context"):
                context_para = document.add_paragraph()
                context_para.add_run("Context: ").italic = True
                context_para.add_run(besluit["context"])
    else:
        _add_paragraph(document, "Er zijn geen formele besluiten genomen in deze vergadering.")
    
    # =========================================================================
    # 7. BESPROKEN ONDERWERPEN
    # =========================================================================
    _add_heading(document, "7. Besproken onderwerpen", level=1)
    
    if analyse.besproken_onderwerpen:
        for i, onderwerp in enumerate(analyse.besproken_onderwerpen, 1):
            titel = onderwerp.get("titel", f"Onderwerp {i}")
            _add_heading(document, f"7.{i} {titel}", level=2)
            
            if onderwerp.get("samenvatting"):
                _add_paragraph(document, onderwerp["samenvatting"])
    else:
        _add_paragraph(document, "Geen gedetailleerde beschrijving van onderwerpen beschikbaar.")
    
    # =========================================================================
    # 8. SPREKERS
    # =========================================================================
    _add_heading(document, "8. Sprekers", level=1)
    
    if analyse.sprekers:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        
        headers = ["Naam", "Functie", "Partij"]
        for i, header in enumerate(headers):
            table.rows[0].cells[i].text = header
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        
        for spreker in analyse.sprekers:
            row = table.add_row()
            row.cells[0].text = spreker.naam or "-"
            row.cells[1].text = spreker.functie or "-"
            row.cells[2].text = spreker.partij or "-"
    else:
        _add_paragraph(document, "Geen sprekersinformatie beschikbaar.")
    
    # =========================================================================
    # 9. OPMERKINGEN
    # =========================================================================
    if analyse.opmerkingen:
        _add_heading(document, "9. Opmerkingen", level=1)
        _add_paragraph(document, analyse.opmerkingen)
    
    # =========================================================================
    # BIJLAGE: TRANSCRIPTIE
    # =========================================================================
    if transcript:
        document.add_page_break()
        _add_heading(document, "Bijlage: Transcriptie", level=1)
        
        # Waarschuwing
        warning = document.add_paragraph()
        warning.add_run(
            "Let op: Deze transcriptie is automatisch gegenereerd en kan fouten bevatten. "
            "Namen en termen kunnen verkeerd zijn getranscribeerd."
        ).italic = True
        
        document.add_paragraph()
        
        # Transcriptie tekst (in kleinere paragrafen voor leesbaarheid)
        paragraphs = transcript.split("\n\n")
        for para in paragraphs:
            if para.strip():
                _add_paragraph(document, para.strip())
    
    document.save(output_path)
    return output_path


def build_teamoverleg_report(
    output_path: Path,
    analyse: "TeamOverlegAnalyse",
    metadata: Optional[Dict[str, str]] = None,
    transcript: Optional[str] = None,
) -> Path:
    """
    Genereer een Word-rapport voor een teamoverleg (bijv. Rotterdams WeerWoord).
    
    Focus op:
    - Samenvatting van de discussie
    - Actielijst met verantwoordelijken en deadlines
    """
    document = Document()
    metadata = metadata or {}
    
    # Titel
    title = document.add_heading(level=0)
    title_run = title.add_run(analyse.titel or "Verslag Teamoverleg")
    title_run.font.size = Pt(18)
    title_run.font.name = "Calibri"
    
    # Subtitel met programma
    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run("Programmateam Rotterdams WeerWoord")
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.italic = True
    
    # Datum
    if analyse.datum or metadata.get("datum"):
        datum_para = document.add_paragraph()
        datum_para.add_run(f"Datum: {analyse.datum or metadata.get('datum', '')}").font.size = Pt(11)
    
    # Automatisch gegenereerd label
    auto_label = document.add_paragraph()
    auto_label.add_run("Dit verslag is automatisch gegenereerd op basis van de opname van het overleg.").italic = True
    
    document.add_paragraph()
    
    # =========================================================================
    # 1. AANWEZIGEN
    # =========================================================================
    if analyse.aanwezigen:
        _add_heading(document, "1. Aanwezigen", level=1)
        aanwezigen_text = ", ".join(analyse.aanwezigen)
        _add_paragraph(document, aanwezigen_text)
    
    # =========================================================================
    # 2. SAMENVATTING
    # =========================================================================
    _add_heading(document, "2. Samenvatting", level=1)
    _add_paragraph(document, analyse.samenvatting or "Geen samenvatting beschikbaar.")
    
    # =========================================================================
    # 3. ACTIELIJST (prominent bovenaan)
    # =========================================================================
    _add_heading(document, "3. Actielijst", level=1)
    
    if analyse.acties:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        
        # Header
        headers = ["Actie", "Verantwoordelijke", "Deadline", "Prioriteit"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
        
        for actie in analyse.acties:
            row = table.add_row()
            row.cells[0].text = actie.actie or "-"
            row.cells[1].text = actie.verantwoordelijke or "-"
            row.cells[2].text = actie.deadline or "-"
            row.cells[3].text = actie.prioriteit or "middel"
        
        document.add_paragraph()
        
        # Telling
        actie_count = document.add_paragraph()
        actie_count.add_run(f"Totaal: {len(analyse.acties)} acties").italic = True
    else:
        _add_paragraph(document, "Er zijn geen specifieke acties afgesproken.")
    
    # =========================================================================
    # 4. BESPROKEN PUNTEN
    # =========================================================================
    _add_heading(document, "4. Besproken punten", level=1)
    
    if analyse.besproken_punten:
        for i, punt in enumerate(analyse.besproken_punten, 1):
            _add_heading(document, f"4.{i} {punt.onderwerp or 'Onderwerp'}", level=2)
            
            if punt.samenvatting:
                _add_paragraph(document, punt.samenvatting)
            
            if punt.besluit:
                besluit_para = document.add_paragraph()
                besluit_para.add_run("Besluit: ").bold = True
                besluit_para.add_run(punt.besluit)
            
            if punt.opmerkingen:
                opmerking_para = document.add_paragraph()
                opmerking_para.add_run("Opmerking: ").italic = True
                opmerking_para.add_run(punt.opmerkingen)
            
            document.add_paragraph()
    else:
        _add_paragraph(document, "Geen gedetailleerde beschrijving van besproken punten beschikbaar.")
    
    # =========================================================================
    # 5. BESLUITEN
    # =========================================================================
    if analyse.besluiten:
        _add_heading(document, "5. Besluiten", level=1)
        for besluit in analyse.besluiten:
            document.add_paragraph(besluit, style="List Bullet")
    
    # =========================================================================
    # 6. AANDACHTSPUNTEN
    # =========================================================================
    if analyse.aandachtspunten:
        _add_heading(document, "6. Aandachtspunten", level=1)
        for punt in analyse.aandachtspunten:
            document.add_paragraph(punt, style="List Bullet")
    
    # =========================================================================
    # 7. VOLGENDE OVERLEG
    # =========================================================================
    if analyse.volgende_overleg:
        _add_heading(document, "7. Volgende overleg", level=1)
        _add_paragraph(document, analyse.volgende_overleg)
    
    # =========================================================================
    # 8. OPMERKINGEN
    # =========================================================================
    if analyse.opmerkingen:
        _add_heading(document, "8. Opmerkingen", level=1)
        _add_paragraph(document, analyse.opmerkingen)
    
    # =========================================================================
    # BIJLAGE: TRANSCRIPTIE
    # =========================================================================
    if transcript:
        document.add_page_break()
        _add_heading(document, "Bijlage: Transcriptie", level=1)
        
        warning = document.add_paragraph()
        warning.add_run(
            "Let op: Deze transcriptie is automatisch gegenereerd en kan fouten bevatten."
        ).italic = True
        
        document.add_paragraph()
        
        paragraphs = transcript.split("\n\n")
        for para in paragraphs:
            if para.strip():
                _add_paragraph(document, para.strip())
    
    document.save(output_path)
    return output_path


def build_gesprek_report(
    output_path: Path,
    analyse: "GesprekAnalyse",
    metadata: Optional[Dict[str, str]] = None,
    transcript: Optional[str] = None,
) -> Path:
    """
    Genereer een Word-rapport voor een algemeen gesprek.

    Werkt voor elk type gesprek: klantgesprek, werkoverleg, 1-op-1, intakegesprek, etc.
    """
    document = Document()
    metadata = metadata or {}

    # Titel
    title = document.add_heading(level=0)
    title_run = title.add_run(analyse.titel or "Verslag gesprek")
    title_run.font.size = Pt(18)
    title_run.font.name = "Calibri"

    if analyse.type_gesprek:
        subtitle = document.add_paragraph()
        subtitle_run = subtitle.add_run(analyse.type_gesprek)
        subtitle_run.font.size = Pt(13)
        subtitle_run.font.italic = True

    auto_label = document.add_paragraph()
    auto_label.add_run("Dit verslag is automatisch gegenereerd op basis van de aangeleverde aantekeningen of opname.").italic = True

    document.add_paragraph()

    # =========================================================================
    # 1. BASISGEGEVENS
    # =========================================================================
    _add_heading(document, "1. Basisgegevens", level=1)

    rows = []
    if analyse.datum or metadata.get("datum"):
        rows.append(("Datum", analyse.datum or metadata.get("datum", "")))
    if analyse.aanwezigen:
        rows.append(("Aanwezigen", ", ".join(analyse.aanwezigen)))
    if analyse.doel:
        rows.append(("Doel", analyse.doel))

    if rows:
        table = document.add_table(rows=len(rows), cols=2)
        table.style = "Table Grid"
        for i, (label, value) in enumerate(rows):
            table.rows[i].cells[0].text = label
            table.rows[i].cells[1].text = value
        document.add_paragraph()
    else:
        _add_paragraph(document, "Geen basisgegevens beschikbaar.")

    # =========================================================================
    # 2. SAMENVATTING
    # =========================================================================
    _add_heading(document, "2. Samenvatting", level=1)
    _add_paragraph(document, analyse.samenvatting or "Geen samenvatting beschikbaar.")

    # =========================================================================
    # 3. BESPROKEN ONDERWERPEN
    # =========================================================================
    _add_heading(document, "3. Besproken onderwerpen", level=1)

    if analyse.onderwerpen:
        for i, onderwerp in enumerate(analyse.onderwerpen, 1):
            _add_heading(document, f"3.{i} {onderwerp.titel or 'Onderwerp'}", level=2)

            if onderwerp.samenvatting:
                _add_paragraph(document, onderwerp.samenvatting)

            if onderwerp.besluit:
                besluit_para = document.add_paragraph()
                besluit_para.add_run("Besluit: ").bold = True
                besluit_para.add_run(onderwerp.besluit)

            document.add_paragraph()
    else:
        _add_paragraph(document, "Geen gedetailleerde onderwerpen beschikbaar.")

    # =========================================================================
    # 4. BESLUITEN
    # =========================================================================
    if analyse.besluiten:
        _add_heading(document, "4. Besluiten", level=1)
        for besluit in analyse.besluiten:
            document.add_paragraph(besluit, style="List Bullet")

    # =========================================================================
    # 5. ACTIELIJST
    # =========================================================================
    _add_heading(document, "5. Actielijst", level=1)

    if analyse.acties:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        headers = ["Actie", "Wie", "Deadline"]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True

        for actie in analyse.acties:
            row = table.add_row()
            row.cells[0].text = actie.actie or "-"
            row.cells[1].text = actie.wie or "-"
            row.cells[2].text = actie.deadline or "-"

        document.add_paragraph()
    else:
        _add_paragraph(document, "Er zijn geen acties afgesproken.")

    # =========================================================================
    # 6. VERVOLGSTAPPEN
    # =========================================================================
    if analyse.vervolgstappen:
        _add_heading(document, "6. Vervolgstappen", level=1)
        for stap in analyse.vervolgstappen:
            document.add_paragraph(stap, style="List Bullet")

    # =========================================================================
    # 7. OPMERKINGEN
    # =========================================================================
    if analyse.opmerkingen:
        _add_heading(document, "7. Opmerkingen", level=1)
        _add_paragraph(document, analyse.opmerkingen)

    # =========================================================================
    # BIJLAGE: ORIGINELE INVOER
    # =========================================================================
    if transcript:
        document.add_page_break()
        _add_heading(document, "Bijlage: Originele invoer", level=1)

        warning = document.add_paragraph()
        warning.add_run(
            "Onderstaande tekst is de originele invoer (aantekeningen of transcriptie) waarop dit verslag is gebaseerd."
        ).italic = True

        document.add_paragraph()

        for para in transcript.split("\n\n"):
            if para.strip():
                _add_paragraph(document, para.strip())

    document.save(output_path)
    return output_path
